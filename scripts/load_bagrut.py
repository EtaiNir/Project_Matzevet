"""
קליטת סבב בגרות: T1 + T2 + המצפן → טבלאות bagrut_* ב-Supabase
(bagrut_students כולל עמודות T2, bagrut_grades, bagrut_subjects, bagrut_tracking).

    python scripts/load_bagrut.py (--accdb <קובץ אקסס> | --xlsx-dir <תיקייה>) [--compass <מצפן>] \
        --authority 800037 --school 800037 --season קיץ --year תשפ"ו [--index] [--dry-run]

--dry-run   מפרק ומדפיס סיכום (ספירות בלבד — בלי שמות, ת"ז או ציונים), לא נוגע במסד.
--index     טוען גם את אינדקס השאלונים הארצי מטבלה 3 באקסס.

מקורות באקסס (מבנה הקובץ של סבא, קיץ תשפ"ו):
    זכאות 1   — פרטי תלמיד + דגלי מעקב
    זכאות 11  — T1 מלל          (קבוצה: חובה)
    זכאות 12  — T1 אנגלית ומתמטיקה
    זכאות 13  — T1 הרחבה, חלק 1 + חלק 2 (פוצל בגלל מגבלת 255 עמודות)
    זכאות 14  — T2 ניתוח AI
    זכאות 3   — אינדקס שאלונים ארצי (עם --index)

פירוק T1 — למה לפי מיקום ולא לפי שם עמודה:
    שמות העמודות נכתבו ביד ואינם עקביים. "MoreshetVeDatziunSofi" (בלי T),
    "Q14375…AravimSofi" (במקום Tziun), "Q8776283…" (ספרה עודפת), ו-
    "Q34211-1/-2/-3" בלי שם בכלל. מה שכן עקבי הוא **הסדר**: כל מקצוע נפתח
    בעמודת ציון סופי ואחריה חמשת מדדיו, ואז שאלונים בשלשות — ציון, משקל,
    משוקלל. לכן הפירוק הולך לפי המיקום, וסמל השאלון של שלשה נקבע ברוב
    קולות בתוכה (כך "8776283" נבלע ב-"776283" של שתי חברותיה).
"""
import argparse
import collections
import glob
import io
import json
import os
import re
import sys

import openpyxl

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
ENV_DB_FILE = os.path.join(HERE, "..", ".env.db")

T_DETAILS = "זכאות 1 - פרטי תתלמידים"
T1_TABLES = [
    ("זכאות 11 - ניתוח רק מקצועות המלל - חדש", "חובה"),
    ("זכאות 12 - ניתוח רק אנגלית ומתמטיקה - חדש", "אנגלית ומתמטיקה"),
    ("זכאות 13 - ניתוח רק מקצועות הרחבה אבו רביע - חלק 1", "מורחב"),
    ("זכאות 13 - ניתוח רק מקצועות הרחבה אבו רביע - חלק 2", "מורחב"),
]
T_T2 = "זכאות 14 - ניתוח AI מהצאט"
T_INDEX = "זכאות 3 - מותאם GPT אינדקס שאלני בגרות של משרד החינוך ללא כפילות"

# מדדי מקצוע ב-T1. הסדר חשוב: "MishkalMitztaber" נבדק לפני "ziunMitztaber".
# שתי שפות כותרת: אקסס ("Anglit3TziunSofi") וקובץ T1 המאוחד ("ציון סופי אנגלית 3 יח"ל").
METRICS = [
    ("final_grade", re.compile(r"(?i)t?ziunsofi$|^ציון סופי ")),
    ("cumulative_weight", re.compile(r"(?i)mishkalmitztaber$|^משקל מצטבר ")),
    ("cumulative_grade", re.compile(r"(?i)t?ziunmitztaber$|^ציון מצטבר ")),
    ("units", re.compile(r"(?i)yechidot$|^יחידות ")),
    ("questionnaires", re.compile(r"(?i)sheelonim$|^שאלונים בוצעו מתוך ")),
    ("completion_status", re.compile(r"(?i)statushashlama$|^סטטוס השלמה ")),
]
# "Q34211…" באקסס, "34211 אזרחות בחינה חיצונית … ציון" בקובץ המאוחד
Q_COL = re.compile(r"^Q?(\d{4,})")
ROLE_BY_SUFFIX = [("weighted", re.compile(r"(?i)meshuklal$|משוקלל$")),
                  ("weight", re.compile(r"(?i)mishkal$|משקל$")),
                  ("grade", re.compile(r"(?i)tziun$|ציון$"))]
POSITIONAL_ROLES = ["grade", "weight", "weighted"]
# מפתח המקצוע = עמודת היחידות בלי הקידומת/סיומת שלה
UNITS_AFFIX = re.compile(r"(?i)yechidot$|^יחידות ")
FINAL_AFFIX = re.compile(r"(?i)t?ziunsofi$|^ציון סופי ")


def alt_groups(pairs):
    """שאלונים שקולים ("או/או") במקצוע: {סמל שאלון: מספר קבוצה}.

    בפגישת 5.10: 150% של זוג שקול אינו טעות — את החלק אפשר להשלים בשאלון
    זה *או* בזה, והתוכנית עצמה תמיד 100%. קבוצה = שאלונים בעלי אותו משקל,
    שספירתם פעם אחת מביאה את המקצוע ל-100% **בדיוק**. בלי התאמה מדויקת —
    לא מסמנים כלום, והאזהרה נשארת (חסר 14% בעברית מורחב הוא בעיה אמיתית).
    """
    import itertools
    weighted = [(c, round(w, 4)) for c, w in pairs if w]
    total = sum(w for _, w in weighted)
    if not weighted or abs(total - 1) < 0.005:
        return {}
    by_w = collections.defaultdict(list)
    for c, w in weighted:
        by_w[w].append(c)
    dups = [w for w, cs in by_w.items() if len(cs) > 1]
    for n in range(1, len(dups) + 1):                  # הקבוצה הקטנה ביותר שמספיקה
        for chosen in itertools.combinations(dups, n):
            collapsed = total - sum(w * (len(by_w[w]) - 1) for w in chosen)
            if abs(collapsed - 1) < 0.005:
                return {c: g for g, w in enumerate(chosen, 1) for c in by_w[w]}
    return {}


def collapsed_weight(pairs):
    """סכום המשקלים כשכל קבוצת שאלונים שקולים נספרת פעם אחת."""
    groups = alt_groups(pairs)
    seen, total = set(), 0.0
    for c, w in pairs:
        g = groups.get(c)
        if g is not None:
            if g in seen:
                continue
            seen.add(g)
        total += w or 0
    return total


# סדר השכבות — לזיהוי השכבה המסיימת (הגבוהה ביותר שיש לה ניתוח זכאות)
GRADE_ORDER = {"ט": 9, "י": 10, "יא": 11, "יב": 12, "יג": 13, "יד": 14}


def grade_rank(g):
    return GRADE_ORDER.get(re.sub(r"[\"'״׳\s]", "", g or ""), -1)


def key_level(key):
    """רמת היחידות מתוך מפתח המקצוע: "Anglit3" או 'אנגלית 3 יח"ל'."""
    m = re.search(r"(\d)$", key) or re.search(r"(\d)\s*יח", key)
    return int(m.group(1)) if m else None

T2_FIELDS = {
    "סטטוס לזכאות לפ T1": "status",
    "מה כבר בוצע": "done_summary",
    "ריכוז מקצועות וציונים לזכאות": "grades_summary",
    "חסמים לזכאות": "blockers",
    "התערבות מומלצת לזכאות": "intervention",
    "אפשרות לזכאות עם שלילי אחד": "one_negative_option",
    "זוג מקצועות לקומפנסציה": "compensation_pair",
    "האם זכאי לשיפוי": "compensation_eligible",
    "מספר שליליים": "negatives_count",
    "סטטוס בשפת אם - ערבית": "mother_tongue_status",
    "סהכ יחידות לתלמיד במקצועות המלל": "core_units",
    "מקצוע מוגבר": "reinforced_subject",
    "סהכ יחדות זכאות": "total_units",
    "תאור סיבת אי הזכאות": "reason",
}
# T2 בקובץ המאוחד (אגיאל) — אותם שדות בניסוח אחר. "האם זכאי לשיפוי" ו"תאור
# סיבת אי הזכאות" אינם בקובץ, ונשארים ריקים.
T2_ALIASES = {
    "מספר זהות": "MsparZehutTalmid",
    "ריכוז ציונים ומקצועות לזכאות": "ריכוז מקצועות וציונים לזכאות",
    "חוסרים וחסמים לזכאות": "חסמים לזכאות",
    "סטטוס זכאות לפי T1": "סטטוס לזכאות לפ T1",
    "התערבות מומלצת": "התערבות מומלצת לזכאות",
    "מספר ציונים שליליים": "מספר שליליים",
    "סטטוס שפת אם ערבית": "סטטוס בשפת אם - ערבית",
    'סה"כ יחידות מלל': "סהכ יחידות לתלמיד במקצועות המלל",
    'סה"כ יחידות זכאות': "סהכ יחדות זכאות",
}

# סוג המקצוע במצפן (עמודת "סוג המקצוע") → קבוצת המקצוע בסבב. בקובץ המאוחד
# כל המקצועות בגיליון אחד, ולכן הקבוצה נגזרת מהמצפן ולא משם הטבלה.
KIND_TO_GROUP = {"מלל": "חובה", "אנגלית": "אנגלית ומתמטיקה", "מתמטיקה": "אנגלית ומתמטיקה",
                 "מורחב": "מורחב"}


# ─────────────────────────────── עזרים ───────────────────────────────

def sid(value):
    """ת"ז כמחרוזת בת 9 ספרות. אקסס שומר אותה כ-Double."""
    if value is None or value == "":
        return None
    return str(int(float(value))).zfill(9)


def txt(value):
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def num(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_table(cur, table, columns=None):
    sel = "*" if columns is None else ", ".join(f"[{c}]" for c in columns)
    cur.execute(f"SELECT {sel} FROM [{table}]")
    names = [d[0] for d in cur.description]
    return names, [dict(zip(names, row)) for row in cur.fetchall()]


# ─────────────────────────────── המצפן ───────────────────────────────

def read_compass(path):
    """שורה לכל מקצוע — סוג, קוד ושם, אחריהם 7 זוגות שאלון/משקל, ואז הערות.
    שורות בלי שם מקצוע (סיכומים) — מדלגים.

    הגיליון והעמודות מזוהים לפי הכותרות בשורה 2 ("שם המקצוע", "שאלון 1"),
    לא לפי שם או מיקום קבוע: באבו רביע ורמות זבולון השם ב-K והשאלונים
    מ-X (והגיליון "גיליון2" או "גיליון4" — ושם "גיליון2" הוא פיבוט ספירות);
    באגיאל השם ב-C והשאלונים מ-Q. הסוג תמיד שתי עמודות לפני השם."""
    wb = openpyxl.load_workbook(path, data_only=True)

    def header_cols(ws):
        heads = {str(ws.cell(2, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
        return heads.get("שם המקצוע"), heads.get("שאלון 1"), heads.get("הערות")

    ws = next((s for s in wb.worksheets if all(header_cols(s)[:2])), None)
    if ws is None:
        sys.exit("לא נמצא במצפן גיליון עם 'שם המקצוע' ו'שאלון 1' בשורה 2")
    name_col, q1_col, notes_col = header_cols(ws)
    notes_col = notes_col or q1_col + 14
    # המקצועות הפנימיים בטבלה תחתונה עם כותרת משלה ("שם המקצוע" שוב). בלי
    # כותרת שנייה — המיקום הישן (שורה 35 ואילך).
    second = next((r for r in range(4, ws.max_row + 1)
                   if str(ws.cell(r, name_col).value or "").strip() == "שם המקצוע"), None)
    internal_from = second or 35
    rows = []
    for r in range(4, ws.max_row + 1):
        name = txt(ws.cell(r, name_col).value)
        if not name or name == "שם המקצוע":
            continue
        pairs = []
        for col in range(q1_col, q1_col + 14, 2):
            code = ws.cell(r, col).value
            if isinstance(code, (int, float)) and code >= 1000:
                pairs.append((int(code), num(ws.cell(r, col + 1).value)))
        if not pairs:
            continue
        kind = txt(ws.cell(r, name_col - 2).value) if name_col > 2 else None  # מלל / אנגלית / מתמטיקה / מורחב
        rows.append({
            "row": r,
            "name": name,
            "kind": kind,
            "internal": r >= internal_from,
            "pairs": pairs,
            "notes": txt(ws.cell(r, notes_col).value),
        })
    return rows


# ─────────────────────────────── T1 ───────────────────────────────

def parse_t1_table(names, rows, group):
    """מחזיר (בלוקים, ציונים, מדדים). בלוק = מקצוע: key, עמודות מדדים, שלשות."""
    blocks = []
    cur_block = None
    pending_q = []

    def flush_q():
        # שלשות לפי הסדר. שארית שאינה שלשה = תקלה במבנה, לא מנחשים.
        if not pending_q:
            return
        if len(pending_q) % 3:
            raise ValueError(f"שאלונים לא בשלשות בבלוק {cur_block['key']}: {pending_q}")
        for i in range(0, len(pending_q), 3):
            trio = pending_q[i:i + 3]
            codes = collections.Counter(int(Q_COL.match(c).group(1)) for c in trio)
            code = codes.most_common(1)[0][0]
            roles = {}
            for pos, col in enumerate(trio):
                role = next((r for r, rx in ROLE_BY_SUFFIX if rx.search(col)), None)
                roles[role or POSITIONAL_ROLES[pos]] = col
            if len(roles) != 3:            # שני "Tziun" בשלשה — חוזרים למיקום
                roles = dict(zip(POSITIONAL_ROLES, trio))
            cur_block["trios"].append((code, roles))
        pending_q.clear()

    for col in names:
        if Q_COL.match(col):
            if cur_block is None:
                raise ValueError(f"עמודת שאלון לפני מקצוע: {col}")
            pending_q.append(col)
            continue
        metric = next((m for m, rx in METRICS if rx.search(col)), None)
        if metric is None:
            continue                                   # עמודות זיהוי
        if metric == "final_grade":
            flush_q()
            cur_block = {"metrics": {}, "trios": [], "group": group, "key": None}
            blocks.append(cur_block)
        elif cur_block is None:
            continue        # "יחידות אנגלית לתלמיד" — עמודת זיהוי לפני המקצוע הראשון
        cur_block["metrics"][metric] = col
        if metric == "units":                          # הקידומת העקבית ביותר
            cur_block["key"] = UNITS_AFFIX.sub("", col).strip()
    flush_q()

    for b in blocks:
        if not b["key"]:
            b["key"] = FINAL_AFFIX.sub("", b["metrics"]["final_grade"]).strip()

    grades, subjects = [], []
    for row in rows:
        student = sid(row.get("MisparZehutChinuch"))
        if not student:
            continue
        for b in blocks:
            m = {k: row.get(c) for k, c in b["metrics"].items()}
            if any(v not in (None, "") for v in m.values()):
                subjects.append({
                    "student_id": student, "subject_key": b["key"],
                    "final_grade": num(m.get("final_grade")),
                    "cumulative_grade": num(m.get("cumulative_grade")),
                    "units": num(m.get("units")),
                    "questionnaires": txt(m.get("questionnaires")),
                    "cumulative_weight": num(m.get("cumulative_weight")),
                    "completion_status": txt(m.get("completion_status")),
                })
            for code, roles in b["trios"]:
                g, w, wd = (num(row.get(roles[r])) for r in POSITIONAL_ROLES)
                if g is None and wd is None:
                    continue
                grades.append({"student_id": student, "subject_key": b["key"],
                               "questionnaire_code": code, "grade": g,
                               "weight": w, "weighted": wd})
    return blocks, grades, subjects


def name_blocks(blocks, compass, index):
    """שם עברי לכל בלוק: המקצוע במצפן שחולק איתו הכי הרבה שאלונים.
    גיבוי: שם המקצוע באינדקס הארצי. אחרון: הקידומת מ-T1."""
    for b in blocks:
        codes = {c for c, _ in b["trios"]}
        best = max(compass, key=lambda s: len(codes & {c for c, _ in s["pairs"]}), default=None)
        if best and codes & {c for c, _ in best["pairs"]}:
            b["name"], b["compass"] = best["name"], best
            continue
        names = collections.Counter(index[c]["subject_name"] for c in codes if c in index and index[c]["subject_name"])
        name = names.most_common(1)[0][0] if names else b["key"]
        # באינדקס "אנגלית" אחת לכל הרמות; הרמה יושבת בקידומת (Anglit3)
        lvl = key_level(b["key"])
        if lvl and not re.search(r"\d", name):
            name = f"{name} {lvl}"
        b["name"], b["compass"] = name, None


# ─────────────────────────────── מקורות ───────────────────────────────
#
# אותם נתונים מגיעים בשתי צורות: קובץ האקסס של סבא (אבו רביע), או ייצוא
# של טבלאות האקסס לקובצי אקסל (רמות זבולון). שני המקורות מחזירים את אותן
# שורות — כך כל הפירוק שאחריהם אחד.

# בייצוא לאקסל, טבלה 1 יצאה עם כותרות בעברית. השמות כאן → השמות באקסס.
DETAIL_ALIASES = {
    "זהות תלמיד": "MisparZehutChinuch", "סמל מוסד": "SemelMosad", "שם המוסד": "SchoolName",
    "פרטי ומשפחה תלמיד": "ShemTalmidChinuch", "שם פרטי": "Prati", "שם משפחה": "Mishpacha",
    "שכבה": "Shichva", "כיתה ושכבה": "KitatEmChinuch", "סטטוס חוסרים": "HaImChaserChinuch",
    "חשד": "Hashad", "עיכוב": "Ikuv", "מאתגר במיוד": "MaatgerimBemyuhad",
    "סיכוי אפסי לזכאות": "ZeroChanceZakaut", "הערות קצר": "HearotTalmidShort",
    "הערות ארוך": "HearotTalmidLong", "שנת לימודים": "ShnatLimud",
}

DETAIL_COLS = ["MisparZehutChinuch", "SemelMosad", "SchoolName", "ShemTalmidChinuch",
               "Prati", "Mishpacha", "Shichva", "KitatEmChinuch", "HaImChaserChinuch",
               "Hashad", "Ikuv", "MaatgerimBemyuhad", "ZeroChanceZakaut",
               "HearotTalmidShort", "HearotTalmidLong"]


def to_bool(v):
    """אקסס נותן True/False; אקסל מחזיר גם 'TRUE', 'FALSE', 0, -1. bool('FALSE') הוא True."""
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v != 0
    return str(v).strip().lower() in ("true", "כן", "yes", "1", "-1", "x", "v")


def list_xlsx(folder):
    """קובצי האקסל בתיקייה — בלי קובצי הנעילה של אקסל ("~$שם.xlsx"), שנוצרים
    כשהקובץ פתוח ונראים כמו הקובץ עצמו לכל זיהוי לפי שם."""
    return sorted(f for f in os.listdir(folder)
                  if f.lower().endswith(".xlsx") and not f.startswith("~$"))


def resolve(path):
    return glob.glob(path)[0] if any(ch in path for ch in "*?") else path


def parse_index(rows):
    index = {}
    for r in rows:
        code = num(r.get("סמל שאלון מערכת 1"))
        if code is None:
            continue
        code = int(code)
        index[code] = {
            "code": code,
            "subject_name": txt(r.get("שם מקצוע מערכת 2")),
            "block": txt(r.get("בלוק 3")),
            "subject_type": txt(r.get("סוג מקצוע 4")),
            "subject_group": txt(r.get("קבוצת המקצוע 5")),
            "exam_form": txt(r.get("צורת היבחנות שאלון 6")),
            "exam_kind": txt(r.get("סוג היבחנות שאלון 7")),
            "units": int(num(r["יחידות מקצוע 13"])) if num(r.get("יחידות מקצוע 13")) is not None else None,
            "weight": num(r.get("משקל יחסי של השאלון 14")),
            "in_final_grade": r.get("משתתף בציון הסופי של המקצוע 15"),
            "required_count": int(num(r["מספר שאלונים נידרש במקצוע 18"])) if num(r.get("מספר שאלונים נידרש במקצוע 18")) is not None else None,
            "relation": txt(r.get("שאלון ראשי או ראשי ותת  תנאי לראשי 19")),
            "notes": txt(r.get("הערות לשאלון")),
        }
    return index


class AccessSource:
    def __init__(self, path):
        import pyodbc          # רק למקור אקסס — בשרת בלי Access Driver קולטים מאקסל
        self.name = os.path.basename(path)
        self.con = pyodbc.connect(r"DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};DBQ="
                                  + path + ";ReadOnly=1")
        self.cur = self.con.cursor()

    def index(self):
        return parse_index(read_table(self.cur, T_INDEX)[1])

    def details(self):
        # בלי עמודת הקבצים המצורפים (Attachment) — pyodbc אינו קורא אותה
        return read_table(self.cur, T_DETAILS, DETAIL_COLS)[1]

    def t1_tables(self):
        for table, group in T1_TABLES:
            names, rows = read_table(self.cur, table)
            yield names, rows, group

    def t2_rows(self):
        return read_table(self.cur, T_T2)[1]

    def close(self):
        self.con.close()


class XlsxSource:
    """תיקייה עם ייצוא האקסל של טבלאות האקסס. מזהה כל קובץ לפי תחילת שמו
    ("זכאות 11 …"), כי סיומת השם משתנה ("אבו רביע" נשאר בשם גם בבית ספר אחר)."""

    T1_PREFIXES = [("זכאות 11", "חובה"), ("זכאות 12", "אנגלית ומתמטיקה"), ("זכאות 13", "מורחב")]

    def __init__(self, folder):
        self.dir = folder
        self.name = os.path.basename(os.path.normpath(folder))
        self.files = list_xlsx(folder)
        self.missing = []

    def _find(self, prefix, many=False):
        # "זכאות 1 " עם רווח — אחרת היא תופסת גם את 11, 12, 13, 14
        hits = [f for f in self.files if f.startswith(prefix)]
        if many:
            return hits
        return hits[0] if hits else None

    def _read(self, fname):
        ws = openpyxl.load_workbook(os.path.join(self.dir, fname), read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        names = [str(h) if h is not None else "" for h in next(it)]
        return names, [dict(zip(names, row)) for row in it if any(v is not None for v in row)]

    def index(self):
        f = self._find("זכאות 3 ")
        return parse_index(self._read(f)[1]) if f else {}

    def details(self):
        f = self._find("זכאות 1 ")
        if not f:
            sys.exit("חסר קובץ פרטי תלמידים (זכאות 1)")
        _, rows = self._read(f)
        return [{DETAIL_ALIASES.get(k, k): v for k, v in r.items()} for r in rows]

    def t1_tables(self):
        for prefix, group in self.T1_PREFIXES:
            files = self._find(prefix, many=True)
            if not files:
                self.missing.append(prefix)      # שגיאה בדוח — לא נטען סבב חלקי בשקט
            for f in files:                       # 13 יכול להגיע בכמה חלקים
                names, rows = self._read(f)
                yield names, rows, group

    def t2_rows(self):
        f = self._find("זכאות 14")
        return self._read(f)[1] if f else []

    def close(self):
        pass


def read_sheet_from_header(path, id_headers):
    """הגיליון הראשון, החל משורת הכותרת — השורה הראשונה (מתוך 10) שיש בה אחת
    מכותרות הת"ז. ב-T2 של אגיאל יש מעליה כותרת ושורה ריקה."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows[:10])
               if any(str(v).strip() in id_headers for v in r if v is not None)), None)
    if hi is None:
        sys.exit(f"לא נמצאה שורת כותרת עם עמודת ת\"ז ב-{os.path.basename(path)}")
    names = [str(h).strip() if h is not None else "" for h in rows[hi]]
    data = [dict(zip(names, r)) for r in rows[hi + 1:] if any(v is not None for v in r)]
    return rows[:hi], names, data


class UnifiedSource:
    """המבנה של אגיאל: קובץ T1 אחד עם כל המקצועות (כותרות בעברית), קובץ T2
    לכל השכבות, ומצפן. אין קובץ פרטי תלמידים — הת"ז, השם והכיתה בעמודות
    הראשונות של T1, והשכבה נגזרת מהכיתה ("יג-3" → "יג"). אין סמל מוסד בקבצים,
    ולכן בדיקת סמל המוסד אינה אפשרית כאן."""

    T1_ID = "מספר זהות חינוך"
    T1_ALIASES = {"מספר זהות חינוך": "MisparZehutChinuch", "שם תלמיד חינוך": "ShemTalmidChinuch",
                  "כיתת אם חינוך": "KitatEmChinuch", "האם חסר חינוך": "HaImChaserChinuch"}

    def __init__(self, folder, t1, t2):
        self.dir = folder
        self.name = os.path.basename(t1)
        self.t1_path = os.path.join(folder, t1)
        self.t2_path = os.path.join(folder, t2) if t2 else None
        self.missing = []        # בלי T2 — אזהרה (no_t2), לא שגיאה
        self.school_name = None
        _, self._t1_names, self._t1_rows = read_sheet_from_header(self.t1_path, {self.T1_ID})
        for r in self._t1_rows:                 # parse_t1_table מחפש את הת"ז בשם האקסס
            r["MisparZehutChinuch"] = r.get(self.T1_ID)

    @staticmethod
    def detect(folder):
        """(t1, t2) אם התיקייה במבנה המאוחד, אחרת None."""
        files = list_xlsx(folder)
        if any(f.startswith("זכאות 11") for f in files):
            return None
        t1 = [f for f in files if re.search(r"(?<![A-Za-z0-9])T1(?![0-9])", f)]
        t2 = [f for f in files if re.search(r"(?<![A-Za-z0-9])T2(?![0-9])", f) or f.startswith("זכאות 14")]
        if len(t1) != 1:
            return None
        return t1[0], (t2[0] if t2 else None)

    def index(self):
        return {}

    def details(self):
        out = []
        for r in self._t1_rows:
            d = {self.T1_ALIASES.get(k, k): v for k, v in r.items() if k in self.T1_ALIASES}
            cls = txt(d.get("KitatEmChinuch"))
            d["Shichva"] = cls.split("-")[0].strip() if cls and "-" in cls else None
            out.append(d)
        return out

    def t1_tables(self):
        # קבוצה None — נקבעת לכל מקצוע לפי הסוג שלו במצפן (KIND_TO_GROUP)
        yield self._t1_names, self._t1_rows, None

    def t2_rows(self):
        if not self.t2_path:
            return []
        above, _, rows = read_sheet_from_header(self.t2_path, {"מספר זהות", "MsparZehutTalmid"})
        # שם בית הספר מכותרת הגיליון: "T2 — ניתוח זכאות … | אגיאל"
        title = next((str(v) for r in above for v in r if v), "")
        if "|" in title:
            self.school_name = title.rsplit("|", 1)[1].strip() or None
        return [{T2_ALIASES.get(k, k): v for k, v in r.items()} for r in rows]

    def close(self):
        pass


def make_source(a):
    if bool(a.accdb) == bool(a.xlsx_dir):
        sys.exit("יש לבחור מקור אחד: --accdb או --xlsx-dir")
    if a.accdb:
        return AccessSource(resolve(a.accdb))
    unified = UnifiedSource.detect(a.xlsx_dir)
    return UnifiedSource(a.xlsx_dir, *unified) if unified else XlsxSource(a.xlsx_dir)


def pg_connect():
    import psycopg2
    load_env_db()
    return psycopg2.connect(host=os.environ["PGHOST"], port=int(os.environ.get("PGPORT", "5432")),
                            dbname=os.environ.get("PGDATABASE", "postgres"),
                            user=os.environ["PGUSER"], password=os.environ["PGPASSWORD"],
                            sslmode="require", connect_timeout=20)


def index_from_db():
    """האינדקס הארצי כבר במסד (נטען עם אבו רביע). קריאה בלבד."""
    pg = pg_connect()
    try:
        with pg.cursor() as c:
            c.execute("select code, subject_name, block, subject_type, subject_group, exam_form, "
                      "exam_kind, units, weight, in_final_grade, required_count, relation, notes "
                      "from public.bagrut_questionnaires")
            cols = [d[0] for d in c.description]
            return {r[0]: dict(zip(cols, r)) for r in c.fetchall()}
    finally:
        pg.close()


# ─────────────────────────────── דוח הבדיקה ───────────────────────────────
#
# המבנה של סבב בגרות קבוע, ולכן אין כאן אישור אדם באמצע (כמו בגנים).
# במקומו — דוח: **שגיאות** עוצרות לפני כל שינוי במסד (נתונים שגויים היו
# נטענים), **אזהרות** נשמרות ומוצגות בהיסטוריה אבל אינן עוצרות (המצפן
# כפי שבית הספר מסר אותו, עם הטעויות שבו — וזה בדיוק מה שהרכז צריך לראות).

def build_report(a, school_name, file_schools, students, grades, subjects, eligibility,
                 blocks, compass, program, tracking, t1_codes, prog_codes,
                 bad_grades, known_ids, src, t2_dropped=None):
    errors, warnings = [], []

    def warn(kind, text, **detail):
        warnings.append({"kind": kind, "text": text, **detail})

    # ── שגיאות: נתונים שגויים היו נכנסים
    if file_schools and file_schools != {str(a.school)}:
        errors.append(f"סמל המוסד בקובץ ({', '.join(sorted(file_schools))}) "
                      f"שונה מהסבב שנבחר ({a.school})")
    if not students:
        errors.append("אין תלמידים בקובץ פרטי התלמידים")
    if not grades or not subjects:
        errors.append("אין ציונים ב-T1")
    for prefix in getattr(src, "missing", []):
        errors.append(f"חסר קובץ {prefix}")

    # ── אזהרות: נטען, ומסומן
    if not compass:
        warn("no_compass", "אין מצפן — המשקלים נגזרו מ-T1")
    for b in blocks:
        pairs = b["compass"]["pairs"] if b.get("compass") else []
        if not pairs:
            continue
        # שקולים נספרים פעם אחת; מקצוע שכולו 0% (בחינות פנימיות) — לא אזהרה
        total = collapsed_weight(pairs)
        if total and abs(total - 1) >= 0.005:
            warn("weight_sum", f"{b['name']}: סכום המשקלים במצפן {round(total * 100)}%",
                 subject=b["name"], percent=round(total * 100))
        # משקל שונה בין המצפן ל-T1 לאותו שאלון
        cw = dict(pairs)
        seen = collections.defaultdict(set)
        for g in grades:
            if g["subject_key"] == b["key"] and g["weight"] is not None:
                seen[g["questionnaire_code"]].add(round(g["weight"], 4))
        for code, ws in seen.items():
            if code in cw and (len(ws) > 1 or abs(next(iter(ws)) - (cw[code] or 0)) >= 0.005):
                warn("weight_mismatch",
                     f"{b['name']}, שאלון {code}: משקל {cw[code]} במצפן מול {sorted(ws)} ב-T1",
                     subject=b["name"], code=code)
    orphans = sorted(t1_codes - prog_codes) if compass else []
    if orphans:
        warn("orphans", f"{len(orphans)} שאלונים עם ציונים שאינם במצפן: "
             + ", ".join(map(str, orphans)), codes=orphans)
    unused = sorted(prog_codes - t1_codes) if compass else []
    if unused:
        warn("unused", f"{len(unused)} שאלונים במצפן בלי אף ציון: "
             + ", ".join(map(str, unused)), codes=unused)
    if bad_grades:
        warn("over_100", f"{bad_grades} ציונים מעל 100")
    unmatched = sum(1 for e in eligibility if e["student_id"] not in known_ids)
    if unmatched:
        warn("t2_unmatched", f"{unmatched} שורות T2 בלי תלמיד בקובץ הפרטים")
    for g, n in sorted((t2_dropped or {}).items()):
        warn("t2_other_grade", f"ניתוח הזכאות כלל גם {n} תלמידי {g} — לא הוצמד להם "
             "(ניתוח הזכאות לשכבה המסיימת בלבד); הם מוצגים לפי התקדמות המקצועות",
             grade=g, count=n)
    if not eligibility:
        warn("no_t2", "אין ניתוח זכאות (T2) — הסבב ייטען עם T1 בלבד")

    checkable = [g for g in grades if None not in (g["grade"], g["weight"], g["weighted"])]
    consistent = sum(1 for g in checkable if abs(g["grade"] * g["weight"] - g["weighted"]) <= 0.51)
    if consistent != len(checkable):
        warn("trio", f"ציון×משקל≠משוקלל ב-{len(checkable) - consistent} שאלונים")

    by_grade = collections.Counter(s["grade"] for s in students)
    return {
        "school_name": school_name,
        "counts": {
            "students": len(students), "by_grade": dict(by_grade), "t2": len(eligibility),
            "grades": len(grades), "subjects": len(blocks), "program": len(program),
            "tracking": len(tracking),
        },
        "subjects": [{"key": b["key"], "name": b["name"], "group": b["group"],
                      "questionnaires": len(b["trios"])} for b in blocks],
        "program_source": "compass" if compass else "derived_from_t1",
        "warnings": warnings,
        "errors": errors,
    }


def write_report(path, report):
    if path:
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(report, fh, ensure_ascii=False, indent=1)


# ─────────────────────────────── ראשי ───────────────────────────────

def load_env_db():
    if not os.path.exists(ENV_DB_FILE):
        return
    for line in open(ENV_DB_FILE, encoding="utf-8"):
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, _, v = line.partition("=")
            os.environ.setdefault(k.strip(), v.strip())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--accdb", help="קובץ האקסס של סבא (T1 + T2 + אינדקס)")
    ap.add_argument("--xlsx-dir", help="תיקייה עם ייצוא האקסל של טבלאות האקסס")
    ap.add_argument("--compass", help="קובץ המצפן. בלעדיו המשקלים נגזרים מ-T1")
    ap.add_argument("--authority", required=True)
    ap.add_argument("--school", required=True)
    ap.add_argument("--season", required=True, choices=["קיץ", "חורף"])
    ap.add_argument("--year", required=True)
    ap.add_argument("--index", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--report-json", help="נתיב לכתיבת דוח הבדיקה (לסוכן ולמסך ההיסטוריה)")
    a = ap.parse_args()

    src = make_source(a)

    # אינדקס ארצי: מהמקור אם יש בו (אקסס של סבא), אחרת מהמסד — שם הוא
    # כבר נטען. משמש לשמות מקצועות כשאין מצפן; נטען למסד רק עם --index.
    index = src.index()
    index_from_source = bool(index)
    if not index:
        try:
            index = index_from_db()
        except KeyError:
            # אין פרטי מסד במחשב הזה. ב-dry-run מספיק המצפן לשמות המקצועות.
            if not a.dry_run:
                raise
            print("(אין פרטי מסד — האינדקס הארצי לא נטען; השמות מהמצפן)")
            index = {}

    compass_path = resolve(a.compass) if a.compass else None
    compass = read_compass(compass_path) if compass_path else []

    students, tracking = [], []
    school_name = None
    file_schools = set()
    for r in src.details():
        s = sid(r.get("MisparZehutChinuch"))
        if not s:
            continue
        school_name = school_name or txt(r.get("SchoolName"))
        sm = num(r.get("SemelMosad"))
        if sm:
            file_schools.add(str(int(sm)))
        students.append({
            "student_id": s, "school_code": a.school,
            "first_name": txt(r.get("Prati")), "last_name": txt(r.get("Mishpacha")),
            "full_name": txt(r.get("ShemTalmidChinuch")),
            "grade": txt(r.get("Shichva")), "class_name": txt(r.get("KitatEmChinuch")),
            "track": None, "missing_status": txt(r.get("HaImChaserChinuch")),
        })
        flags = dict(suspected=to_bool(r.get("Hashad")), has_blocker=to_bool(r.get("Ikuv")),
                     fighting=to_bool(r.get("MaatgerimBemyuhad")),
                     zero_chance=to_bool(r.get("ZeroChanceZakaut")),
                     note_short=txt(r.get("HearotTalmidShort")), note_long=txt(r.get("HearotTalmidLong")))
        if any(flags.values()):
            tracking.append({"school_code": a.school, "student_id": s, **flags})

    all_blocks, grades, subjects = [], [], []
    for names, rows, group in src.t1_tables():
        b, g, sb = parse_t1_table(names, rows, group)
        all_blocks += b
        grades += g
        subjects += sb

    name_blocks(all_blocks, compass, index)
    # קובץ T1 מאוחד: הקבוצה לפי סוג המקצוע במצפן
    for b in all_blocks:
        if b["group"] is None:
            c = b["compass"]
            b["group"] = ("פנימי" if c and c["internal"]
                          else KIND_TO_GROUP.get(c["kind"], "מורחב") if c else "מורחב")

    # מפתח כפול = אותו מקצוע פוצל בין שתי טבלאות — שגיאה במבנה, לא ממזגים בשקט.
    dup = [k for k, n in collections.Counter(b["key"] for b in all_blocks).items() if n > 1]
    if dup:
        sys.exit(f"מקצוע מופיע ביותר מבלוק אחד: {dup}")

    eligibility = []
    for r in src.t2_rows():
        s = sid(r.get("MsparZehutTalmid"))
        if not s:
            continue
        rec = {"student_id": s}
        for src_col, dst in T2_FIELDS.items():
            v = r.get(src_col)
            rec[dst] = (int(num(v)) if num(v) is not None else None) if dst == "negatives_count" else                        (num(v) if dst == "core_units" else txt(v))
        # שורה בלי סטטוס = תלמיד שלא נותח (בייצוא מאקסל מגיעות גם שורות י"ב ריקות)
        if rec["status"] is None:
            continue
        eligibility.append(rec)
    school_name = school_name or getattr(src, "school_name", None)
    src.close()

    # ── ניתוח הזכאות — לשכבה המסיימת בלבד (פגישת 5.10: "הפרדה גמורה בין
    # י"ב ל-י"ג"). המסכים מזהים את השכבה המסיימת לפי מי שיש לו T2; קובץ T2
    # שכולל גם את י"ב (אגיאל) היה מערבב את שתי השכבות במסך י"ג. התלמידים
    # עצמם נשארים — רק ניתוח הזכאות אינו מוצמד לשכבה שאינה מסיימת.
    grade_of = {s["student_id"]: s["grade"] for s in students}
    t2_grades = {grade_of[e["student_id"]] for e in eligibility if grade_of.get(e["student_id"])}
    t2_dropped = collections.Counter()
    if len(t2_grades) > 1:
        top = max(t2_grades, key=grade_rank)
        kept = []
        for e in eligibility:
            g = grade_of.get(e["student_id"])
            if g and g != top:
                t2_dropped[g] += 1
            else:
                kept.append(e)
        eligibility = kept

    # ── המצפן → מקצועות הסבב ותוכנית. מקצוע במצפן שאין לו בלוק ב-T1
    # (למשל הפנימיים) נכנס עם מפתח משלו, כדי שהתוכנית תהיה שלמה.
    by_compass_row = {id(b["compass"]): b for b in all_blocks if b["compass"]}
    round_subjects, program = [], []
    for i, b in enumerate(all_blocks):
        units = collections.Counter(index[c]["units"] for c, _ in b["trios"]
                                    if c in index and index[c]["units"]).most_common(1)
        # גיבוי אחרון: היחידות שב-T1 עצמו (עמודת "יחידות …" של המקצוע)
        units = units or collections.Counter(int(sb["units"]) for sb in subjects
                                             if sb["subject_key"] == b["key"] and sb["units"]).most_common(1)
        lvl = key_level(b["key"])
        round_subjects.append({"subject_key": b["key"], "subject_name": b["name"],
                               "subject_group": b["group"],
                               "units": lvl or (units[0][0] if units else None),
                               "sort": i})
    if not compass:
        # אין מצפן: המשקל של כל שאלון נגזר מ-T1 עצמו — לכל ציון מצורף
        # המשקל שלו. הערך השכיח לכל (מקצוע, שאלון) נכנס לתוכנית, ומסומן.
        for b in all_blocks:
            for k, (code, _) in enumerate(b["trios"]):
                ws = collections.Counter(round(g["weight"], 4) for g in grades
                                         if g["subject_key"] == b["key"] and g["questionnaire_code"] == code
                                         and g["weight"] is not None)
                if not ws and not any(g["questionnaire_code"] == code and g["subject_key"] == b["key"] for g in grades):
                    continue        # שאלון שאיש לא ניגש אליו — אין ממה לגזור
                program.append({"subject_key": b["key"], "questionnaire_code": code,
                                "weight": ws.most_common(1)[0][0] if ws else None, "sort": k,
                                "notes": "נגזר מ-T1 — אין מצפן לסבב" if k == 0 else None,
                                "alt_group": None})
    for j, s in enumerate(compass):
        b = by_compass_row.get(id(s))
        key = b["key"] if b else f"P{s['row']}"
        if not b:
            round_subjects.append({"subject_key": key, "subject_name": s["name"],
                                   "subject_group": "פנימי" if s["internal"] else "מורחב",
                                   "units": None, "sort": 1000 + j})
        groups = alt_groups(s["pairs"])
        for k, (code, weight) in enumerate(s["pairs"]):
            program.append({"subject_key": key, "questionnaire_code": code, "weight": weight,
                            "sort": k, "notes": s["notes"] if k == 0 else None,
                            "alt_group": groups.get(code)})

    # ── סיכום (ספירות בלבד)
    grade_by = collections.Counter(s["grade"] for s in students)
    graduating = collections.Counter(
        next((s["grade"] for s in students if s["student_id"] == e["student_id"]), None)
        for e in eligibility).most_common(1)
    graduating = graduating[0][0] if graduating else None
    t1_codes = {g["questionnaire_code"] for g in grades}
    prog_codes = {p["questionnaire_code"] for p in program}
    consistent = sum(1 for g in grades if None not in (g["grade"], g["weight"], g["weighted"])
                     and abs(g["grade"] * g["weight"] - g["weighted"]) <= 0.51)
    checkable = sum(1 for g in grades if None not in (g["grade"], g["weight"], g["weighted"]))
    bad_grades = sum(1 for g in grades if g["grade"] is not None and g["grade"] > 100)
    known_ids = {s["student_id"] for s in students}

    print(f"בית ספר: {school_name} ({a.school}) · {a.season} {a.year} · רשות {a.authority}")
    print(f"תלמידים: {len(students)} · לפי שכבה: {dict(grade_by)} · שכבה מסיימת (T2): {graduating}")
    print(f"מקצועות ב-T1: {len(all_blocks)} · ציוני שאלון: {len(grades)} · שורות מדדים: {len(subjects)}")
    for b in all_blocks:
        print(f"   {b['group']:<16} {b['key']:<34} → {b['name']:<28} שאלונים: {len(b['trios'])}")
    print(f"בדיקת שלשות: ציון×משקל=משוקלל ב-{consistent}/{checkable}")
    print(f"ציון מעל 100: {bad_grades}")
    print(f"T2: {len(eligibility)} · מתוכם בלי פרטי תלמיד: {sum(1 for e in eligibility if e['student_id'] not in known_ids)}")
    print(f"מצפן: {len(compass)} מקצועות, {len(program)} שאלונים" + ("" if compass else " — נגזר מ-T1, אין קובץ מצפן"))
    print(f"   שאלונים ב-T1 שאינם במצפן: {len(t1_codes - prog_codes)} · במצפן ואין להם ציון: {len(prog_codes - t1_codes)}")
    print(f"מעקב — תלמידים עם דגל או הערה: {len(tracking)}")
    print(f"אינדקס ארצי: {len(index)} שאלונים ({'מהמקור' if index_from_source else 'מהמסד'})")

    report = build_report(a, school_name, file_schools, students, grades, subjects, eligibility,
                          all_blocks, compass, program, tracking, t1_codes, prog_codes,
                          bad_grades, known_ids, src, t2_dropped)
    for w in report["warnings"]:
        print(f"⚠ {w['text']}")
    for e in report["errors"]:
        print(f"✗ {e}")
    write_report(a.report_json, report)
    if report["errors"]:
        sys.exit("הקליטה נעצרה לפני כל שינוי במסד: " + " · ".join(report["errors"]))

    if a.dry_run:
        print("\n(--dry-run — המסד לא נגע)")
        return

    # מלכודת 30: טעינה שמוחקת ואז מכניסה מסרבת על קלט ריק *לפני* המחיקה.
    if not students or not grades or not subjects:
        sys.exit("קלט ריק — לא טוען (ולא מוחק את הסבב הקיים)")

    import psycopg2
    from psycopg2.extras import execute_values, Json
    load_env_db()
    pg = psycopg2.connect(host=os.environ["PGHOST"], port=int(os.environ.get("PGPORT", "5432")),
                          dbname=os.environ.get("PGDATABASE", "postgres"),
                          user=os.environ["PGUSER"], password=os.environ["PGPASSWORD"],
                          sslmode="require", connect_timeout=20)
    pg.autocommit = False
    code = a.authority
    try:
        with pg.cursor() as c:
            c.execute("select 1 from public.authorities where code = %s", (code,))
            if not c.fetchone():
                sys.exit(f"הרשות {code} אינה קיימת — יש ליצור אותה במסך הניהול קודם")
            if a.index and index_from_source:
                cols = list(next(iter(index.values())).keys())
                execute_values(c, f"""
                    insert into public.bagrut_questionnaires ({', '.join(cols)}) values %s
                    on conflict (code) do update set
                    {', '.join(f'{k} = excluded.{k}' for k in cols if k != 'code')}, updated_at = now()""",
                    [tuple(q[k] for k in cols) for q in index.values()])

            # סבב קיים עם אותו מפתח נמחק על כל שורותיו (cascade). המעקב נשאר —
            # הוא לפי בית ספר + ת"ז, לא לפי סבב.
            c.execute("""delete from public.bagrut_rounds
                         where authority_code=%s and school_code=%s and season=%s and school_year=%s""",
                      (code, a.school, a.season, a.year))
            stats = {"program_source": "compass" if compass else "derived_from_t1",
                     "students": len(students), "grades": len(grades), "subjects": len(subjects),
                     "eligibility": len(eligibility), "program": len(program),
                     "orphans_t1": len(t1_codes - prog_codes), "grades_over_100": bad_grades}
            c.execute("""insert into public.bagrut_rounds
                           (authority_code, school_code, school_name, season, school_year,
                            graduating_grade, source, stats)
                         values (%s,%s,%s,%s,%s,%s,%s,%s) returning id""",
                      (code, a.school, school_name, a.season, a.year, graduating,
                       src.name, Json(stats)))
            rid = c.fetchone()[0]

            def put(table, rows, cols):
                execute_values(c, f"insert into public.{table} (round_id, {', '.join(cols)}) values %s",
                               [(rid, *[r[k] for k in cols]) for r in rows], page_size=1000)

            put("bagrut_round_subjects", round_subjects,
                ["subject_key", "subject_name", "subject_group", "units", "sort"])
            put("bagrut_program", program,
                ["subject_key", "questionnaire_code", "weight", "sort", "notes", "alt_group"])
            # T2 נכנס לשורת התלמיד. תלמיד בלי T2 מקבל עמודות ריקות ו-in_t2=false.
            t2_by_id = {e["student_id"]: e for e in eligibility}
            t2_cols = list(T2_FIELDS.values())
            for s in students:
                e = t2_by_id.get(s["student_id"])
                s["authority_code"] = code
                s["in_t2"] = e is not None
                for k in t2_cols:
                    s[k] = e[k] if e else None
            put("bagrut_students", students,
                ["student_id", "authority_code", "school_code", "first_name", "last_name",
                 "full_name", "grade", "class_name", "track", "missing_status", "in_t2"] + t2_cols)
            put("bagrut_grades", grades,
                ["student_id", "subject_key", "questionnaire_code", "grade", "weight", "weighted"])
            put("bagrut_subjects", subjects,
                ["student_id", "subject_key", "final_grade", "cumulative_grade", "units",
                 "questionnaires", "cumulative_weight", "completion_status"])

            # מעקב: רק מה שעוד לא קיים. עריכה שנעשתה באתר גוברת על אקסס.
            if tracking:
                tcols = ["authority_code", "school_code", "student_id", "suspected", "has_blocker",
                         "fighting", "zero_chance", "note_short", "note_long"]
                execute_values(c, f"""insert into public.bagrut_tracking ({', '.join(tcols)})
                                      values %s on conflict (authority_code, school_code, student_id)
                                      do nothing""",
                               [tuple({**t, "authority_code": code}[k] for k in tcols) for t in tracking])
        pg.commit()
        report["round_id"] = str(rid)
        write_report(a.report_json, report)
        print(f"\n✓ נטען. סבב {rid}")
    except Exception:
        pg.rollback()
        raise
    finally:
        pg.close()


if __name__ == "__main__":
    main()
