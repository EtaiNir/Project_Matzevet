"""
קליטת קובץ דרג ב' (גנים, לידה עד 3, קידום נוער, חינוך ביתי) לטבלת
students_{code}_tier_b.

    python scripts/load_tier_b.py --file "גנים.xls" --code 1400000 --dry-run

═══════════════════════════════════════════════════════════════════
העיקרון שמפריד בין ה-AI לקוד
═══════════════════════════════════════════════════════════════════
ה-AI מחליט **איזו עמודה היא מה** — ותו לא. כל ההמרות (ריפוד ת"ז, ספרת
הביקורת של סמל הישוב, תאריכי אקסל, ניקוי) נעשות כאן בקוד, דטרמיניסטית.
מודל אינו אמור לנסח כלל כזה, ובוודאי לא לנסח אותו אחרת בכל הרצה.

ומה שנשלח אליו: **שורת הכותרת ופרופיל צורה לכל עמודה** — לא ערכים.
עמודה שערכיה נראים כת"ז, טלפון, מייל או שם מקבלת תיאור ("9 ספרות")
במקום דוגמה. אף נתון של ילד אינו יוצא מהבניין.

המיפוי שאושר נשמר לפי טביעת אצבע של שורת הכותרת, ולכן קובץ באותן
כותרות נקלט בחודש הבא **בלי שום קריאת API**.

═══════════════════════════════════════════════════════════════════
שתי מלכודות שהקובץ הזה נבנה סביבן
═══════════════════════════════════════════════════════════════════
1. **סמל ישוב מקוצר מול ארוך.** קובץ הגנים של מנשה מגיע עם 213
   (משמרות), בעוד שהטבלה הראשית מחזיקה 2139. ה-RLS משווה מול
   scope_values שנבנו מהקוד הארוך, ולכן טעינה בלי המרה הייתה גורמת
   לכך שמנהל בהיקף יישובי יראה **אפס ילדי גן** — בלי שגיאה ובלי רמז.

2. **ת"ז נשמרת בקובץ כמספר.** אפסים מובילים כבר נמחקו במקור: ת.ז. הורה
   מופיעה כ-29422003 ומשמעה 029422003. בלי ריפוד ל-9 זיהוי האחים מול
   הטבלה הראשית פשוט לא היה תופס.
"""
import argparse
import hashlib
import json
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

import psycopg2
from psycopg2.extras import execute_values

sys.path.insert(0, str(Path(__file__).resolve().parent))
import tier_b_fields as F  # noqa: E402

# הפלט חייב להיות UTF-8 בלי תלות בקודפייג של הקונסולה.
#
# ב-Windows ברירת המחדל היא cp1255, ואז ה-JSON שהסוכן קורא ב---emit-mapping
# חוזר בקידוד אחר — `json.loads` נופל על
# `'utf-8' codec can't decode byte 0xf9`. זה עבד במקרה רק מפני שהסוכן
# מגדיר PYTHONIOENCODING לתהליך הבן; סקריפט שנכון רק כשהקורא שלו זוכר
# להגדיר משתנה סביבה אינו נכון. לכן ההגדרה כאן, במקור.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):   # פלט שאינו טקסט (pipe/קובץ)
        pass

DASHBOARD_DIR = Path(__file__).resolve().parent.parent
BACKUP_DIR = DASHBOARD_DIR / "agent" / "backups"
AI_MODEL = "claude-sonnet-5"

# טבלת סמלי הישוב הארצית: קוד קצר (כפי שמגיע בקובצי רשות) → קוד ארוך
# (כפי שבמצב"ת). יושבת עם שאר המילונים של ה-pipeline. ראה LocalityResolver.
LOCALITY_TABLE = DASHBOARD_DIR.parent / "Itay_Modules" / "tables" / "locality_codes.json"

NULL_TOKENS = {"", "nan", "none", "nat", "null", "-", "--"}

# שלב החינוך שנכתב לכל קבוצה של דרג ב'.
#
# לגנים נבחר **"גן ילדים בלבד"** ולא "גנים": זה הניסוח בקובץ מוסדות
# משרד החינוך, והוא עקבי עם המשמעות של העמודה בטבלה הראשית — שם היא
# מתארת את שלבי ה*מוסד* ("יסודי בלבד", "חט"ב + עליונה") ולא את הילד.
# לשלוש האחרות אין מקבילה במשרד החינוך, ולכן שם הקבוצה הוא גם השלב.
#
# הקבוצה עצמה נשמרת בנפרד ב-`source_group`, שמוצג בבורר השדות בשם
# "קבוצת אוכלוסייה" — כך אפשר לסנן לפי קבוצה בלי להסתמך על הניסוח כאן.
STAGE_BY_GROUP = {
    "גנים": "גן ילדים בלבד",
    "לידה עד 3": "לידה עד 3",
    "קידום נוער": "קידום נוער",
    "חינוך ביתי": "חינוך ביתי",
}

# ערך מיוחד במיפוי: "צור עמודה תוספתית חדשה בשם הכותרת".
# ⚠️ חייב להיות זהה ל-NEW_COLUMN ב-src/lib/admin.ts.
NEW_COLUMN = "__new__"

# הקבוצות שבהן ההבחנה "משובץ" מול "מועמד" נגזרת משיבוץ למוסד.
#
# בגנים ההבחנה אמיתית: יש ילד עם גן, ויש ילד שממתין לשיבוץ. בקידום נוער,
# בחינוך ביתי ובלידה עד 3 אין מוסד כלל — ילד שנמצא ברשימה **נמצא
# בתוכנית**. בלי הכלל הזה כל הקבוצה הייתה נטענת כ"מועמד", וסינון ברירת
# המחדל במסך («מצב רישום = משובץ») היה מסתיר אותה כולה — הנתונים במסד,
# המסך ריק, ואף שגיאה.
PLACEMENT_DECIDES_STATUS = {"גנים"}


# ═══════════════════════════ סביבה ═══════════════════════════

def load_env() -> None:
    """טוען .env.db ו-.env.agent, בלי לדרוס משתנים שכבר הוגדרו."""
    for name in (".env.db", ".env.agent"):
        path = DASHBOARD_DIR / name
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8-sig").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip("\"'"))


def connect():
    return psycopg2.connect(
        host=os.environ["PGHOST"],
        port=int(os.environ.get("PGPORT", "5432")),
        dbname=os.environ.get("PGDATABASE", "postgres"),
        user=os.environ["PGUSER"],
        password=os.environ["PGPASSWORD"],
        sslmode="require",
        connect_timeout=20,
    )


# ═══════════════════════ קריאת הקובץ ═══════════════════════

def read_table(path: Path) -> tuple[list[str], list[dict]]:
    """
    מחזיר (כותרות, שורות). תומך ב-.xls ישן וב-.xlsx.

    שורת הכותרת אינה בהכרח הראשונה — קבצים מרשויות מגיעים לפעמים עם
    שורת כותרת עליונה או שורה ריקה. נבחרת השורה עם הכי הרבה תאי טקסט
    לא-ריקים מבין 12 הראשונות.
    """
    suffix = path.suffix.lower()
    if suffix == ".xls":
        grid = _read_xls(path)
    elif suffix in (".xlsx", ".xlsm"):
        grid = _read_xlsx(path)
    else:
        raise SystemExit(f"סוג קובץ לא נתמך: {suffix} (נדרש .xls או .xlsx)")

    if not grid:
        raise SystemExit("הקובץ ריק")

    best, best_n = 0, -1
    for r in range(min(12, len(grid))):
        n = sum(1 for v in grid[r] if isinstance(v, str) and v.strip())
        if n > best_n:
            best, best_n = r, n

    raw_headers = grid[best]
    headers, seen = [], {}
    for i, h in enumerate(raw_headers):
        name = str(h).strip() if h is not None else ""
        if not name:
            name = f"עמודה_{i + 1}"
        # כותרת כפולה תדרוס את קודמתה בדיקט — ממספרים אותה
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        headers.append(name)

    rows = []
    for r in range(best + 1, len(grid)):
        values = grid[r]
        if not any(str(v).strip() for v in values if v is not None):
            continue  # שורה ריקה לגמרי
        rows.append({headers[c]: (values[c] if c < len(values) else None)
                     for c in range(len(headers))})
    return headers, rows


def _read_xls(path: Path) -> list[list]:
    import xlrd
    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_index(0)
    grid = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            ctype, value = sheet.cell_type(r, c), sheet.cell_value(r, c)
            if ctype == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                except Exception:
                    value = None
            elif ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                # מספר שלם באקסל הוא float — ‎233492149.0 אינו ת"ז תקינה
                value = int(value)
            elif ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                value = None
            row.append(value)
        grid.append(row)
    return grid


def _read_xlsx(path: Path) -> list[list]:
    from openpyxl import load_workbook
    book = load_workbook(str(path), data_only=True, read_only=True)
    sheet = book[book.sheetnames[0]]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


# ═════════════════ פרופיל עמודות — בלי נתונים אישיים ═════════════════

ID_RE = re.compile(r"^\d{7,9}$")
PHONE_RE = re.compile(r"^0\d{1,2}-?\d{7}$|^0\d{8,9}$")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _looks_sensitive(header: str, values: list) -> bool:
    # "שם ישוב" ו"שם מוסד" נתפסים ברמז "שם", אבל אין בהם אדם — ודווקא
    # הערכים שלהם הם מה שמאפשר ל-AI לזהות את העמודה. הנייטרליות גוברת.
    if any(hint in header for hint in F.NEUTRAL_HEADER_HINTS):
        return False
    if any(hint in header for hint in F.SENSITIVE_HEADER_HINTS):
        return True
    sample = [str(v).strip() for v in values[:40] if str(v).strip()]
    if not sample:
        return False
    hits = sum(1 for v in sample
               if ID_RE.match(v) or PHONE_RE.match(v) or EMAIL_RE.match(v))
    return hits >= len(sample) * 0.6


def build_profiles(headers: list[str], rows: list[dict]) -> list[dict]:
    """
    פרופיל לכל עמודה — זה מה שנשלח ל-AI.

    עמודה אישית מתוארת בצורתה ולא בערכיה. עמודה נייטרלית (שנת לימודים,
    שם ישוב, שם מוסד) כן מקבלת דוגמאות אמיתיות, כי הן באמת עוזרות
    למיפוי ואין בהן איש.
    """
    profiles = []
    for h in headers:
        values = [r[h] for r in rows if r.get(h) is not None
                  and str(r[h]).strip() not in NULL_TOKENS]
        filled = len(values)
        profile = {"כותרת": h, "מלאים": filled, "מתוך": len(rows)}

        if filled == 0:
            profile["דוגמה"] = "ריקה לחלוטין"
            profiles.append(profile)
            continue

        first = values[0]
        if isinstance(first, (datetime, date)):
            profile["סוג"] = "תאריך"
            profile["דוגמה"] = "תאריך"
        elif isinstance(first, (int, float)):
            profile["סוג"] = "מספר"
            lengths = sorted({len(str(v)) for v in values[:200]})
            profile["אורך"] = lengths[0] if len(lengths) == 1 else lengths
        else:
            profile["סוג"] = "טקסט"

        if _looks_sensitive(h, values):
            lengths = sorted({len(str(v).strip()) for v in values[:200]})
            shape = f"{lengths[0]} תווים" if len(lengths) == 1 else f"{lengths[0]}–{lengths[-1]} תווים"
            if any(EMAIL_RE.match(str(v).strip()) for v in values[:20]):
                shape = "כתובת דוא\"ל"
            elif str(first).strip().startswith("0"):
                shape += ", מתחיל ב-0"
            profile["דוגמה"] = shape
            profile["אישי"] = True
        else:
            seen, samples = set(), []
            for v in values:
                s = str(v).strip()
                if isinstance(v, (datetime, date)):
                    s = v.strftime("%d/%m/%Y")
                if s and s not in seen and len(s) <= 24:
                    seen.add(s)
                    samples.append(s)
                if len(samples) == 3:
                    break
            profile["דוגמה"] = ", ".join(samples)
        profiles.append(profile)
    return profiles


# ═══════════════════════ מיפוי הכותרות ═══════════════════════

def headers_hash(headers: list[str]) -> str:
    norm = [re.sub(r"\s+", " ", h).strip() for h in headers]
    return hashlib.sha256(json.dumps(norm, ensure_ascii=False).encode()).hexdigest()[:32]


def saved_mapping(conn, code: str, group: str, h_hash: str) -> dict | None:
    with conn.cursor() as cur:
        cur.execute(
            """select mapping from public.tier_b_mappings
                where authority_code = %s and source_group = %s and headers_hash = %s""",
            (code, group, h_hash))
        row = cur.fetchone()
    return row[0] if row else None


def save_mapping(conn, code: str, group: str, h_hash: str,
                 headers: list[str], mapping: dict) -> None:
    with conn.cursor() as cur:
        cur.execute(
            """insert into public.tier_b_mappings
                 (authority_code, source_group, headers_hash, headers, mapping, use_count, used_at)
               values (%s, %s, %s, %s, %s, 1, now())
               on conflict (authority_code, source_group, headers_hash)
               do update set mapping   = excluded.mapping,
                             used_at   = now(),
                             use_count = public.tier_b_mappings.use_count + 1""",
            (code, group, h_hash, headers, json.dumps(mapping, ensure_ascii=False)))
    # ⚠️ בלי ה-commit הזה השמירה אבדה בשקט: `load()` סוגר את הטרנזקציה
    # שלו, ההכנסה כאן נשארה פתוחה, ו-`conn.close()` ב-finally גלגל אותה
    # אחורה — בעוד המסך הדפיס "המיפוי נשמר". הודעת הצלחה על פעולה שלא
    # קרתה היא גרועה מכישלון גלוי.
    conn.commit()


def _claude(prompt: str, max_tokens: int = 4000) -> str:
    """קריאה אחת ל-Claude. מחזיר את הטקסט של התשובה."""
    import requests

    key = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if not key:
        raise RuntimeError("אין ANTHROPIC_API_KEY ב-.env.agent")

    response = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers={"x-api-key": key,
                 "anthropic-version": "2023-06-01",
                 "content-type": "application/json"},
        json={"model": AI_MODEL, "max_tokens": max_tokens,
              "messages": [{"role": "user", "content": prompt}]},
        timeout=180)

    if response.status_code != 200:
        detail = response.text[:300]
        if "credit balance" in detail:
            raise RuntimeError(
                "חשבון ה-Anthropic ללא יתרה. אפשר להוסיף קרדיט ב-Plans & Billing, "
                "או להריץ עם ‎--mapping לקובץ מיפוי ידני.")
        raise RuntimeError(f"קריאת ה-AI נכשלה ({response.status_code}): {detail}")

    # לא להניח ש-content[0] הוא טקסט. בפרומפט מורכב המודל חושב קודם, והתשובה
    # מגיעה כשני בלוקים: thinking ואחריו text. `content[0]["text"]` נפל על
    # `KeyError: 'text'` — כשל שנראה כמו שגיאת API, בעוד שהסטטוס הוא 200.
    blocks = response.json().get("content", [])
    text = "".join(b.get("text", "") for b in blocks if b.get("type") == "text").strip()
    if not text:
        kinds = ", ".join(b.get("type", "?") for b in blocks) or "ריק"
        raise RuntimeError(
            f"ה-AI לא החזיר טקסט. סוגי הבלוקים שהתקבלו: {kinds}. "
            "אפשר להריץ עם ‎--mapping לקובץ מיפוי ידני.")
    return text


class AIResponseError(RuntimeError):
    """התשובה הגיעה, אבל אינה JSON תקין. שונה משגיאת API — כאן יש טעם לנסות שוב."""


def _parse_json_object(text: str) -> dict:
    """מחלץ אובייקט JSON מהתשובה — גם אם נוספה לו גדר קוד או משפט לפניו."""
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.M).strip()
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < start:
        raise AIResponseError(f"ה-AI לא החזיר JSON: {text[:200]}")
    try:
        return json.loads(text[start:end + 1])
    except json.JSONDecodeError as exc:
        raise AIResponseError(f"ה-AI לא החזיר JSON תקין: {exc}\n{text[:300]}")


def ai_propose(profiles: list[dict], targets: list[dict], note: str = "") -> dict:
    """
    סבב מיפוי מול רשימת יעד נתונה.

    מחזיר { כותרת: {"field": מפתח או None, "confidence": "high" | "low"} }.
    נשלחת סכימה בלבד — ראה build_profiles.
    """
    prompt = (
        "אתה ממפה עמודות של קובץ אקסל שהתקבל מרשות מקומית, לשדות של מסד "
        "נתוני תלמידים.\n"
        f"{note}\n\n"
        f"העמודות בקובץ:\n{json.dumps(profiles, ensure_ascii=False, indent=1)}\n\n"
        f"שדות היעד האפשריים:\n{json.dumps(targets, ensure_ascii=False, indent=1)}\n\n"
        "החזר JSON בלבד, בלי הסבר ובלי גדר קוד, בצורה:\n"
        '{"<כותרת בקובץ>": {"field": "<מפתח שדה היעד או null>", '
        '"confidence": "high" או "low"}}\n\n'
        "כללים:\n"
        "- כל כותרת חייבת להופיע, גם אם field הוא null.\n"
        "- field=null לעמודה שאין לה שדה יעד מתאים ברשימה, או שהיא ריקה לחלוטין.\n"
        "- אסור למפות שתי כותרות לאותו שדה יעד.\n"
        "- confidence=high רק כשיש ברשימה שדה אחד בלבד שמתאים בבירור.\n"
        "- confidence=low כשיש כמה שדות סבירים, כשהכותרת עמומה, או כשהסוג לא תואם.\n"
        "- שים לב להבדל בין פרטי הילד לבין פרטי ההורים: «נייד 2 תלמיד» ו«נייד 2 "
        "הורה 1» הם שדות שונים.\n"
        "- 'תז' של הילד היא מספר הזהות שלו; 'תז אב'/'תז אם' הן של הורה 1/2.\n"
        "- אל תנחש: כשלא ברור, null."
    )

    # מודל מחזיר לפעמים JSON פגום — נמדד: פסיק חסר בשתיים מתוך תשע קריאות.
    # בלי ניסיון חוזר, תשובה פגומה אחת מכשילה את הקליטה כולה. חוזרים רק על
    # פגם בתשובה; שגיאת API או יתרה עולות מיד — ניסיון חוזר לא יעזור שם.
    attempts = 3
    for attempt in range(1, attempts + 1):
        try:
            raw = _parse_json_object(_claude(prompt))
            break
        except AIResponseError as exc:
            if attempt == attempts:
                raise
            print(f"  ⚠ תשובה פגומה מה-AI (ניסיון {attempt}/{attempts}) — מנסה שוב: "
                  f"{str(exc).splitlines()[0][:80]}")

    out = {}
    for header, value in raw.items():
        if isinstance(value, dict):
            field, confidence = value.get("field"), value.get("confidence")
        else:                       # תשובה בפורמט הישן: כותרת → שדה
            field, confidence = value, "low"
        out[header] = {"field": field or None,
                       "confidence": "high" if confidence == "high" else "low"}
    return out


def ai_mapping(profiles: list[dict]) -> tuple[dict, dict]:
    """
    מיפוי בקריאה אחת, מול **כל** שדות היעד.

    ═══════════════════════════════════════════════════════════════
    למה רשימה אחת ולא שתי שכבות
    ═══════════════════════════════════════════════════════════════
    נוסתה גם חלוקה לשתיים — קודם ~25 שדות מרכזיים, ורק השאריות מול
    הרשימה המלאה. היא **הפסידה במדידה**: 83/90 מול 90/90 לרשימה השטוחה,
    ו-126/133 מול 132/133 על כל הקבצים.

    הסיבה מבנית ולא אקראית: הקריאה הראשונה רואה רק את המרכזיים, ולכן
    בוחרת "שדה שבערך מתאים" במקום להשאיר null — והעמודה כבר אינה מגיעה
    לקריאה השנייה. «טלפון קווי אם» נכנס ל«טלפון בית» של הילד בכל שלוש
    ההרצות. כשכל השדות מוצגים יחד, המודל מבחין ביניהם.

    מחזיר (מיפוי, מטא). המטא הוא { כותרת: {"tier", "confidence"} }:
    `tier` הוא "core" לשדה מרכזי (שיש לו המרה מיוחדת) ו-"extended" לכל
    שדה אחר, ולפיו מסך האישור מסמן בירוק או בצהוב.
    """
    # עמודה ריקה לגמרי לא ממופה — בקוד, לא בהנחיה. נמדד שהמודל מיפה עמודה
    # ריקה לשדה **בביטחון גבוה** למרות ההנחיה המפורשת; כלל דטרמיניסטי אינו
    # תלוי בשכנוע.
    empty = {p["כותרת"] for p in profiles if p.get("מלאים", 0) == 0}
    askable = [p for p in profiles if p["כותרת"] not in empty]

    mapping = {p["כותרת"]: None for p in profiles}
    meta: dict[str, dict] = {}
    if not askable:
        return mapping, meta

    proposed = ai_propose(askable, F.all_targets_for_prompt())
    used: set[str] = set()
    for profile in askable:
        header = profile["כותרת"]
        result = proposed.get(header) or {}
        field = result.get("field")
        if field in F.CATALOG and field not in used:
            mapping[header] = field
            used.add(field)
            meta[header] = {
                "tier": "core" if field in F.TARGETS else "extended",
                "confidence": result.get("confidence", "low"),
            }
    return mapping, meta


def validate_mapping(mapping: dict, headers: list[str],
                     allowed: dict | set | None = None) -> dict:
    """
    מנקה מיפוי: משאיר רק כותרות שקיימות ושדות יעד חוקיים, בלי כפילויות.

    שדה יעד חוקי = אחד מהשדות המוכרים, **או** כל עמודה שקיימת בטבלה
    (`allowed`) — המשתמש יכול לבחור במסך האישור כל אחת מהן. שדות שהטעינה
    או הדפדפן כותבים בעצמם (`RESERVED`) נדחים תמיד.
    """
    clean, used = {}, {}
    for header, target in (mapping or {}).items():
        if header not in headers or not target:
            continue
        if target == NEW_COLUMN:
            # אינה שדה בטבלה, ולכן גם אינה כפופה לכלל "שדה אחד לכל כותרת":
            # כמה כותרות יכולות לבקש עמודה תוספתית חדשה, כל אחת משלה.
            clean[header] = target
            continue
        if target in F.RESERVED:
            print(f"  ⚠ שדה שהמערכת כותבת בעצמה, מדולג: {target}")
            continue
        if target not in F.TARGETS and not (allowed and target in allowed):
            print(f"  ⚠ שדה יעד שאינו בטבלה, מדולג: {target}")
            continue
        if target in used:
            print(f"  ⚠ '{target}' כבר ממופה מ-'{used[target]}' — '{header}' מדולג")
            continue
        used[target] = header
        clean[header] = target
    return clean


# ═══════════════════════ המרות ═══════════════════════

def luhn_check_digit(payload: str) -> str:
    """ספרת הביקורת שמוסיפים לסמל ישוב מקוצר כדי לקבל את הארוך."""
    total = 0
    for i, ch in enumerate(reversed(payload)):
        d = int(ch)
        if i % 2 == 0:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return str((10 - total % 10) % 10)


def valid_israeli_id(value: str) -> bool:
    if not re.fullmatch(r"\d{9}", value):
        return False
    total = 0
    for i, ch in enumerate(value):
        d = int(ch) * (1 if i % 2 == 0 else 2)
        total += d if d < 10 else d - 9
    return total % 10 == 0


def as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


class LocalityResolver:
    """
    ממיר סמל ישוב לקוד הארוך שבמצב"ת — לפי **טבלה**, לא לפי חישוב.

    הקוד הארוך הוא בדרך כלל הקצר + ספרת ביקורת Luhn, אבל לא תמיד: 6 יישובים
    מאוחדים קיבלו קוד חדש שאינו נגזר (בנימינה-גבעת עדה 9800 → 13524, ולא
    98004). החישוב שמר אותם שגויים בשקט.

    ויש מלכודת שנייה: 84 ערכים הם גם קוד קצר של ישוב אחד וגם קוד ארוך של
    ישוב אחר. הוספת ספרה "ליתר ביטחון" לערך שכבר ארוך הייתה משייכת ילד
    לישוב אחר. לכן לכל קובץ נקבעת **מוסכמה אחת** — קצר או ארוך — לפי רוב
    הערכים החד-משמעיים בו, והיא מכריעה את הדו-משמעיים.
    """

    def __init__(self, known_main: set[str], table_path: Path = LOCALITY_TABLE):
        self.known_main = known_main
        self.short_to_long: dict[str, str] = {}
        self.long_codes: set[str] = set()
        self.has_table = table_path.exists()
        if self.has_table:
            data = json.loads(table_path.read_text(encoding="utf-8"))
            for row in data["localities"]:
                self.short_to_long[row["short"]] = row["long"]
                self.long_codes.add(row["long"])
        self.convention = "short"

    def _classify(self, digits: str) -> tuple[bool, bool]:
        is_short = digits in self.short_to_long
        is_long = digits in self.long_codes or digits in self.known_main
        return is_short, is_long

    def decide(self, values) -> tuple[int, int]:
        """קובע אם הקובץ כתוב בקודים קצרים או ארוכים. מחזיר (קצרים, ארוכים)."""
        short = long_ = 0
        for value in values:
            digits = re.sub(r"\D", "", as_text(value))
            if not digits:
                continue
            is_short, is_long = self._classify(digits)
            if is_short and not is_long:
                short += 1
            elif is_long and not is_short:
                long_ += 1
        self.convention = "long" if long_ > short else "short"
        return short, long_

    def resolve(self, digits: str, warn: list) -> str:
        is_short, is_long = self._classify(digits)
        if is_short and is_long:
            return self.short_to_long[digits] if self.convention == "short" else digits
        if is_short:
            return self.short_to_long[digits]
        if is_long:
            return digits
        # לא בטבלה הארצית ולא במצב"ת. ספרת ביקורת רק אם התוצאה מאומתת
        # במצב"ת של הרשות — אחרת שומרים כמות שהוא, בלי להמציא קוד.
        candidate = digits + luhn_check_digit(digits)
        if candidate in self.known_main:
            return candidate
        warn.append(f"סמל ישוב {digits} אינו בטבלה הארצית — נשמר כמות שהוא")
        return digits


def convert(value, kind: str, localities: "LocalityResolver", warn: list) -> object:
    text = as_text(value)
    if text.lower() in NULL_TOKENS:
        return None

    if kind == "id":
        digits = re.sub(r"\D", "", text)
        if not digits:
            return None
        return digits.zfill(9)[-9:] if len(digits) <= 9 else digits[:9]

    if kind == "locality":
        digits = re.sub(r"\D", "", text)
        if not digits:
            return None
        return localities.resolve(digits, warn)

    if kind == "mosad":
        digits = re.sub(r"\D", "", text)
        return digits.zfill(6) if digits else None

    if kind == "date":
        if isinstance(value, (datetime, date)):
            return value.date() if isinstance(value, datetime) else value
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return None

    if kind == "house_no":
        # 0 הוא מציין "אין מספר בית", לא ערך. תא ריק במקור נשאר ריק.
        return None if text in ("0", "0.0") else text

    if kind == "phone":
        cleaned = re.sub(r"[^\d-]", "", text)
        return cleaned or None

    if kind == "gender":
        first = text[:1]
        if first in ("ז", "1", "M", "m"):
            return "זכר"
        if first in ("נ", "2", "F", "f"):
            return "נקבה"
        return text or None

    if kind == "year":
        digits = re.sub(r"\D", "", text)
        return digits or None

    if kind == "number":
        # עמודה מספרית במסד: ערך שאינו מספר היה מפיל את כל הטעינה
        # (Postgres דוחה את ה-insert כולו), ולכן הוא הופך לריק.
        cleaned = text.replace(",", "")
        try:
            float(cleaned)
            return cleaned
        except ValueError:
            return None

    return text or None


def transform(rows: list[dict], mapping: dict, group: str,
              localities: "LocalityResolver",
              kinds: dict | None = None) -> tuple[list[dict], list[dict]]:
    """
    מחיל את המיפוי וההמרות. מחזיר (שורות נקיות, שורות שנדחו).

    `kinds` — סוג ההמרה לכל שדה יעד, כולל עמודות שאינן ברשימת השדות
    המוכרים (ראה `F.kind_for`). בלעדיו נופלים לסוג המוכר או לטקסט.
    """
    clean, rejected = [], []
    kinds = kinds or {}

    for index, raw in enumerate(rows, start=2):   # 2 = השורה הראשונה אחרי הכותרת
        out, warn, extra = {}, [], {}
        for header, target in mapping.items():
            if target == NEW_COLUMN:
                # לא נכנס לשורה של הטבלה — נכתב אחר כך לעמודות התוספתיות,
                # תחת המפתח של הכותרת. ראה load_extra_values.
                value = as_text(raw.get(header))
                if value and value.lower() not in NULL_TOKENS:
                    extra[header] = value
                continue
            out[target] = convert(raw.get(header), kinds.get(target) or F.kind(target),
                                  localities, warn)

        missing = [F.label(f) for f in F.REQUIRED if not out.get(f)]
        if missing:
            rejected.append({"שורה": index, "סיבה": "חסר " + ", ".join(missing)})
            continue

        student_id = out["MISPAR_ZEHUT"]
        if not valid_israeli_id(student_id):
            # לא דוחים — סבא: "עדיף להשאיר תלמיד במערכת מאשר למחוק אותו"
            warn.append("ת\"ז אינה עוברת בדיקת ספרת ביקורת")

        # ── שדות נגזרים ──
        out["source_group"] = group
        # שלב החינוך נכתב **לכל** הקבוצות, לא רק לגנים. בגרסה הראשונה
        # היה כאן `if group == "גנים"`, ושלוש הקבוצות האחרות היו נטענות
        # עם עמודת שלב חינוך ריקה — בלי שאיש ישים לב עד שיסננו לפיה.
        out.setdefault("SHLAV_HINUCH_TEUR", STAGE_BY_GROUP.get(group, group))

        # ילד עם שיבוץ נחשב "משובץ" — זה מה שסינון ברירת המחדל במסך
        # מחפש (ACTIVE_STATUS_FIELD ב-presets.ts). בלי זה הוא נטען
        # לטבלה ופשוט לא מופיע.
        # אם הקובץ עצמו מיפה מצב רישום — הוא גובר; אחרת נגזר לפי הקבוצה.
        if not out.get("MATZAV_RISHUM_TEUR"):
            out["MATZAV_RISHUM_TEUR"] = (
                ("משובץ" if out.get("SEMEL_MOSAD") else "מועמד")
                if group in PLACEMENT_DECIDES_STATUS else "משובץ")

        if extra:
            out["_extra"] = extra
        if warn:
            out["_warn"] = warn
        clean.append(out)

    # כפילויות ת"ז בתוך הקובץ עצמו — האחרונה מנצחת, כמו ב-load_main
    by_id, duplicates = {}, 0
    for row in clean:
        if row["MISPAR_ZEHUT"] in by_id:
            duplicates += 1
        by_id[row["MISPAR_ZEHUT"]] = row
    if duplicates:
        print(f"  ⚠ {duplicates} כפילויות ת\"ז בתוך הקובץ — נשמרה האחרונה")

    return list(by_id.values()), rejected


# ═══════════════════════ מסד הנתונים ═══════════════════════

def known_locality_codes(conn, code: str) -> set[str]:
    with conn.cursor() as cur:
        cur.execute(f'select distinct "SEMEL_YISHUV1" from public.students_{code} '
                    'where "SEMEL_YISHUV1" is not null')
        return {str(r[0]) for r in cur.fetchall()}


def table_columns(conn, table: str) -> dict[str, str]:
    """עמודות הטבלה וסוגן — הסוג קובע את ההמרה לעמודה שאינה מוכרת מראש."""
    with conn.cursor() as cur:
        cur.execute("""select column_name, data_type from information_schema.columns
                        where table_schema = 'public' and table_name = %s""", (table,))
        return {r[0]: r[1] for r in cur.fetchall()}


def backup_existing(conn, code: str, group: str) -> Path | None:
    """
    גיבוי לפני דריסה — אותו עיקרון כמו בעדכון המצב"ת: קליטה חדשה מוחקת
    את הקבוצה הקיימת, ובלי עותק אין דרך לחזור.
    """
    table = f"students_{code}_tier_b"
    with conn.cursor() as cur:
        cur.execute(f'select count(*) from public.{table} where source_group = %s', (group,))
        if cur.fetchone()[0] == 0:
            return None
        cur.execute(f'select * from public.{table} where source_group = %s', (group,))
        columns = [d[0] for d in cur.description]
        rows = cur.fetchall()

    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = group[:30]
    sheet.sheet_view.rightToLeft = True
    sheet.append(columns)
    for row in rows:
        sheet.append(["" if v is None else str(v) for v in row])

    folder = BACKUP_DIR / table
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{table}_{group}_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    book.save(path)
    return path


def ensure_extra_column(conn, code: str, label: str) -> str:
    """
    מזהה העמודה התוספתית בשם הזה, או יוצר אותה.

    ההתאמה לפי **תווית** ולא לפי מזהה, כדי שקליטה בחודש הבא תכתוב לאותה
    עמודה במקום ליצור כפילות. `view_id is null` — עמודה של הרשות כולה,
    ולא של טבלה ייעודית אחת.
    """
    with conn.cursor() as cur:
        cur.execute("""select id from public.extra_columns
                        where authority_code = %s and view_id is null and label = %s
                        limit 1""", (code, label))
        row = cur.fetchone()
        if row:
            return str(row[0])
        cur.execute("""insert into public.extra_columns
                         (authority_code, label, type, sort_order)
                       values (%s, %s, 'text',
                               coalesce((select max(sort_order) + 1
                                           from public.extra_columns
                                          where authority_code = %s), 0))
                       returning id""", (code, label, code))
        return str(cur.fetchone()[0])


def load_extra_values(conn, code: str, rows: list[dict]) -> tuple[int, int]:
    """
    כותב עמודות שאין להן שדה בטבלה אל המנגנון של "עמודות שהמשתמש מוסיף":
    קטלוג `extra_columns` + jsonb ב-`students_{code}_extra`.

    זה בדיוק מה שהמנגנון ההוא נועד לו. קובץ של קידום נוער יביא «גורם
    מטפל» ו«סטטוס נשירה», ולקובץ של לידה עד 3 יהיה «שם מטפלת» — שדות
    שאינם קיימים בסכימת משרד החינוך ואין להם לאן להיכנס. כאן הם נשמרים
    בלי לשנות סכימה, בטבלה נפרדת שהעדכון החודשי אינו נוגע בה.

    ⚠️ הטבלה התוספתית ממופתחת ב**ת"ז בלבד**, בלי קבוצה. ילד שנמצא בשתי
    קבוצות של דרג ב' מחזיק ערכים תוספתיים משותפים לשתיהן — הוא אותו ילד.

    מחזיר (כמה עמודות, כמה ילדים).
    """
    labels = {label for row in rows for label in (row.get("_extra") or {})}
    if not labels:
        return 0, 0

    with conn.cursor() as cur:
        cur.execute("select public._ensure_extra_table(%s)", (code,))
    ids = {label: ensure_extra_column(conn, code, label) for label in sorted(labels)}

    payload = []
    for row in rows:
        extra = row.get("_extra") or {}
        if not extra:
            continue
        data = {ids[label]: value for label, value in extra.items()}
        payload.append((row["MISPAR_ZEHUT"], json.dumps(data, ensure_ascii=False)))

    if payload:
        table = f"students_{code}_extra"
        with conn.cursor() as cur:
            # מיזוג ולא דריסה: `||` שומר ערכים שהוזנו ידנית בעמודות אחרות.
            execute_values(cur, f'''
                insert into public.{table} ("MISPAR_ZEHUT", data, updated_at)
                values %s
                on conflict ("MISPAR_ZEHUT") do update
                   set data = public.{table}.data || excluded.data,
                       updated_at = now()
            ''', payload, template="(%s, %s::jsonb, now())", page_size=500)
        conn.commit()
    return len(ids), len(payload)


def load(conn, code: str, group: str, rows: list[dict]) -> int:
    """דריסה סלקטיבית: רק הקבוצה הזו, רק ברשות הזו, בטרנזקציה אחת."""
    table = f"students_{code}_tier_b"
    columns = sorted({k for row in rows for k in row if not k.startswith("_")})
    quoted = ", ".join(f'"{c}"' for c in columns)
    values = [tuple(row.get(c) for c in columns) for row in rows]

    with conn.cursor() as cur:
        cur.execute(f'delete from public.{table} where source_group = %s', (group,))
        execute_values(cur,
                       f'insert into public.{table} ({quoted}) values %s',
                       values, page_size=500)
        cur.execute(f'select count(*) from public.{table} where source_group = %s', (group,))
        total = cur.fetchone()[0]
    conn.commit()
    return total


# ═══════════════════════ main ═══════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description="קליטת קובץ דרג ב'")
    parser.add_argument("--file", required=True)
    parser.add_argument("--code", required=True)
    parser.add_argument("--group", default="גנים")
    parser.add_argument("--dry-run", action="store_true",
                        help="לנתח, למפות ולדווח — בלי לגעת במסד")
    parser.add_argument("--mapping", help="קובץ JSON של מיפוי ידני")
    parser.add_argument("--no-ai", action="store_true")
    parser.add_argument("--save-mapping", action="store_true",
                        help="לשמור את המיפוי לשימוש חוזר")
    parser.add_argument("--emit-mapping", action="store_true",
                        help="לפתור את המיפוי, להדפיס JSON ולצאת — בלי לטעון. "
                             "זה מה שהסוכן מריץ בשלב הראשון, לפני אישור אדם.")
    args = parser.parse_args()

    if not args.code.isdigit():
        raise SystemExit("קוד רשות חייב להכיל ספרות בלבד")

    # ב-emit-mapping הפלט התקני הוא JSON שהסוכן קורא, ולכן כל ההודעות
    # לאדם עוברות ל-stderr. בלי זה שורת לוג אחת הייתה שוברת את הפענוח.
    _real_stdout = sys.stdout
    if args.emit_mapping:
        sys.stdout = sys.stderr

    load_env()
    path = Path(args.file)
    if not path.exists():
        raise SystemExit(f"הקובץ לא נמצא: {path}")

    print(f"קורא {path.name}…")
    headers, rows = read_table(path)
    print(f"  {len(rows)} שורות, {len(headers)} עמודות")

    profiles = build_profiles(headers, rows)
    h_hash = headers_hash(headers)

    conn = connect()
    try:
        columns = table_columns(conn, f"students_{args.code}_tier_b")
        if not columns and not args.emit_mapping:
            raise SystemExit(
                f"הטבלה students_{args.code}_tier_b אינה קיימת — "
                "יש להריץ קודם את מיגרציה 023")

        # ── מיפוי: שמור → ידני → AI ──
        source, meta = None, {}
        if args.mapping:
            mapping = json.loads(Path(args.mapping).read_text(encoding="utf-8"))
            source = "manual"
        else:
            mapping = saved_mapping(conn, args.code, args.group, h_hash)
            source = "saved" if mapping else None
            if not mapping and not args.no_ai:
                print("  אין מיפוי שמור לכותרות האלה — פונה ל-AI…")
                mapping, meta = ai_mapping(profiles)
                source = "ai"
            elif not mapping:
                raise SystemExit(
                    "אין מיפוי שמור, ו-‎--no-ai הוגדר. יש לספק ‎--mapping.")

        raw_mapping = dict(mapping or {})
        mapping = validate_mapping(raw_mapping, headers, columns)

        # מיפוי שאדם אישר (manual — כך מגיע האישור מהמסך) חייב להיטען
        # כמו שאושר. אם עמודה שבחר נפלה בוולידציה, עדיף לסרב ברעש מאשר
        # לטעון בלעדיה ולתת לו לגלות אחר כך שהיא חסרה.
        dropped = {h: t for h, t in raw_mapping.items()
                   if t and h in headers and h not in mapping}
        if dropped and source == "manual":
            raise SystemExit(
                "חלק מהמיפוי שאושר אינו תקין ולא ייטען: " +
                ", ".join(f"{h} → {t}" for h, t in dropped.items()) +
                ". לא נגעתי בנתונים הקיימים.")

        print(f"\nמיפוי ({source}) — {len(mapping)} עמודות:")
        for header, target in mapping.items():
            print(f"   {header:<26} →  {F.label(target)}  [{target}]")
        unmapped = [h for h in headers if h not in mapping]
        if unmapped:
            print(f"   לא ממופות: {', '.join(unmapped)}")

        if args.emit_mapping:
            # כל כותרת מופיעה, גם כזו שלא מופתה — מסך האישור מציג את
            # השורה ומאפשר לבחור לה שדה. עמודה שנעלמת מהרשימה נראית
            # כאילו המערכת לא ראתה אותה.
            sys.stdout = _real_stdout
            print(json.dumps({
                "mapping": {h: mapping.get(h) for h in headers},
                "source": source,
                "meta": {h: m for h, m in meta.items() if mapping.get(h)},
                "headers": headers,
                "rows": len(rows),
            }, ensure_ascii=False))
            return

        missing_cols = [t for t in mapping.values()
                        if t != NEW_COLUMN and t not in columns]
        if missing_cols:
            raise SystemExit(f"שדות שאינם בטבלה: {missing_cols}")

        # ── המרה ──
        known = known_locality_codes(conn, args.code)
        kinds = {t: F.kind_for(t, columns.get(t)) for t in mapping.values()}

        localities = LocalityResolver(known)
        locality_headers = [h for h, t in mapping.items() if kinds.get(t) == "locality"]
        if locality_headers:
            if not localities.has_table:
                print(f"  ⚠ הטבלה הארצית לא נמצאה ({LOCALITY_TABLE}) — "
                      "סמלי ישוב ייבדקו רק מול המצב\"ת")
            short, long_ = localities.decide(
                r.get(h) for r in rows for h in locality_headers)
            print("  סמלי ישוב בקובץ: " +
                  ("קצרים — יומרו לארוכים" if localities.convention == "short"
                   else "ארוכים — נשמרים כמות שהם") +
                  f" (קצרים {short} · ארוכים {long_})")

        clean, rejected = transform(rows, mapping, args.group, localities, kinds)

        warned = [r for r in clean if r.get("_warn")]
        print(f"\nתוצאה: {len(clean)} נקלטות · {len(rejected)} נדחו · "
              f"{len(warned)} עם אזהרה")
        for r in rejected[:10]:
            print(f"   ✗ שורה {r['שורה']}: {r['סיבה']}")
        if len(rejected) > 10:
            print(f"   … ועוד {len(rejected) - 10}")

        seen_warnings = {}
        for row in warned:
            for w in row["_warn"]:
                seen_warnings[w] = seen_warnings.get(w, 0) + 1
        for text, count in sorted(seen_warnings.items(), key=lambda x: -x[1])[:8]:
            print(f"   ⚠ {text}  ({count})")

        placed = sum(1 for r in clean if r.get("MATZAV_RISHUM_TEUR") == "משובץ")
        print(f"\n   משובצים: {placed} · מועמדים: {len(clean) - placed}")
        new_columns = sorted({h for h, t in mapping.items() if t == NEW_COLUMN})
        if new_columns:
            print(f"   עמודות תוספתיות שייווצרו/יתעדכנו: {', '.join(new_columns)}")

        # ⚠️ מחסום לפני כל כתיבה. הטעינה מוחקת את כל הקבוצה ואז מכניסה —
        # ולכן מיפוי שבור (ריק, או בלי ת"ז) היה מוחק את כל ילדי הגן ומכניס
        # אפס, בשקט, כשהגיבוי הוא רשת הביטחון היחידה. עדיף לסרב ברעש.
        unmapped_required = [F.label(f) for f in F.REQUIRED if f not in mapping.values()]
        if unmapped_required:
            raise SystemExit(
                "המיפוי חסר שדות חובה: " + ", ".join(unmapped_required) +
                ". לא נגעתי בנתונים הקיימים.")
        if not clean:
            raise SystemExit(
                f"אף שורה לא עברה ({len(rejected)} נדחו). "
                "סירבתי לטעון — טעינה של אפס שורות הייתה מוחקת את הקבוצה כולה. "
                "לא נגעתי בנתונים הקיימים.")

        if args.dry_run:
            print("\n[--dry-run] לא נגעתי במסד.")
            return

        backup = backup_existing(conn, args.code, args.group)
        if backup:
            print(f"\nגיבוי: {backup.name}")

        total = load(conn, args.code, args.group, clean)
        print(f"נטענו {len(clean)} שורות. בטבלה כעת {total} בקבוצה '{args.group}'.")

        cols, kids = load_extra_values(conn, args.code, clean)
        if cols:
            print(f"עמודות תוספתיות: {cols} · ערכים ל-{kids:,} ילדים")

        if args.save_mapping and source in ("ai", "manual"):
            save_mapping(conn, args.code, args.group, h_hash, headers, mapping)
            print("המיפוי נשמר — הקובץ הבא באותן כותרות ייקלט בלי AI.")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
